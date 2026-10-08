"""10. API assessment creation  11. analyst override  + end-to-end pipeline and export."""
from __future__ import annotations

import io
import json

from openpyxl import load_workbook

from app.db.session import SessionLocal
from app.models.entities import Assessment
from app.services.llm.factory import set_llm_provider
from app.services.presentation import REDUCED_MODE_NOTICE
from tests.conftest import FakeLLM


def _create(client, **kw):
    # these tests analyse the same sample company repeatedly: each run is an explicit new analysis
    body = {"company_name": "Nordvik Components AB", "year": 2025, "enable_web_search": False, "confirm_new": True, **kw}
    r = client.post("/api/assessments", json=body)
    assert r.status_code == 201, r.text
    return r.json()


def _upload(client, aid, pdf):
    with open(pdf, "rb") as f:
        r = client.post(f"/api/assessments/{aid}/documents", files={"file": ("../../evil name?.pdf", f, "application/pdf")})
    assert r.status_code == 201, r.text
    return r.json()


def _run(client, aid):
    r = client.post(f"/api/assessments/{aid}/run")
    assert r.status_code == 202, r.text
    st = client.get(f"/api/assessments/{aid}/status").json()  # TestClient runs background tasks synchronously
    assert st["status"] == "COMPLETED", st
    return st


def _matrix(client, aid):
    recs = client.get(f"/api/assessments/{aid}/recommendations").json()
    return {(r["factor"], r["dimension"]): r for r in recs}


def test_app_starts_without_model(client):
    assert client.get("/api/health").json() == {"status": "ok"}
    st = client.get("/api/llm/status").json()
    assert st["available"] is False
    assert "Local LLM model is not configured" in st["message"]
    q = client.get("/api/config/questionnaire").json()
    assert len(q["factors"]) == 5 and q["dimensions"]["planning"]["max_level"] == 6


def test_create_assessment_and_validation(client):
    a = _create(client, website="example.com", country="Sweden", industry="Manufacturing")
    assert a["status"] == "CREATED" and a["company"]["website"] == "https://example.com" and a["company"]["domain"] == "example.com"
    assert client.post("/api/assessments", json={"company_name": " ", "year": 2025}).status_code == 422
    assert client.post("/api/assessments", json={"company_name": "X", "year": 1800}).status_code == 422
    assert client.post("/api/assessments", json={"company_name": "X", "year": 2025, "website": "http://127.0.0.1"}).status_code == 422
    assert client.get("/api/assessments/99999").status_code == 404
    audit = client.get(f"/api/assessments/{a['id']}/audit").json()
    assert audit[0]["event_type"] == "assessment_created"


def test_upload_validation(client, tmp_path):
    a = _create(client)
    txt = tmp_path / "notes.txt"
    txt.write_text("hello")
    with open(txt, "rb") as f:
        assert client.post(f"/api/assessments/{a['id']}/documents", files={"file": ("notes.txt", f, "text/plain")}).status_code == 415
    fake = tmp_path / "fake.pdf"
    fake.write_text("not really a pdf")
    with open(fake, "rb") as f:
        assert client.post(f"/api/assessments/{a['id']}/documents", files={"file": ("fake.pdf", f, "application/pdf")}).status_code == 415


def test_full_pipeline_heuristic_mode(client, sample_pdf):
    a = _create(client)
    src = _upload(client, a["id"], sample_pdf)
    assert src["original_filename"] == "evil name_.pdf" and src["source_type"] == "uploaded_report"
    st = _run(client, a["id"])
    # the technical reason is kept internally; analysts only see a plain notice
    assert st["notices"] == [REDUCED_MODE_NOTICE]
    with SessionLocal() as db:
        assert any("Local LLM model is not configured" in w for w in db.get(Assessment, a["id"]).warnings)

    m = _matrix(client, a["id"])
    assert len(m) == 1 + 5 * 3
    assert m[("reporting", "reporting")]["recommended_level"] == 6
    assert m[("energy_management", "planning")]["recommended_level"] == 5
    assert m[("energy_management", "execution")]["recommended_level"] == 3
    assert m[("energy_management", "performance")]["recommended_level"] == 2
    assert m[("diversity_inclusion", "planning")]["recommended_level"] == 5
    assert m[("diversity_inclusion", "execution")]["recommended_level"] == 3  # "in line with our diversity strategy"
    assert m[("diversity_inclusion", "performance")]["recommended_level"] == 2  # strategy wording is not a sector benchmark
    assert m[("labor_practices", "execution")]["recommended_level"] == 2  # activity not explicitly tied to a goal
    assert m[("product_lifecycle", "planning")]["recommended_level"] == 4
    assert m[("data_security", "planning")]["recommended_level"] == 4
    # data gaps are "not found", never "does not exist"
    gap = m[("data_security", "performance")]
    assert gap["recommended_level"] == 1 and gap["evidence_status"] == "NO_EVIDENCE_FOUND" and gap["evidence_ids"] == []
    assert gap["reason"].startswith("No evidence of")
    labor_perf = m[("labor_practices", "performance")]
    assert labor_perf["recommended_level"] == 1 and any("unquantified" in n for n in labor_perf["notes"])
    # heuristic mode never yields silently-accepted results
    for r in m.values():
        if r["evidence_ids"]:
            assert r["evidence_status"] != "FOUND"
            assert r["evidence_ids"]

    # every recommendation above level 1 is backed by verified evidence with page + quote
    detail = client.get(f"/api/recommendations/{m[('energy_management', 'planning')]['id']}").json()
    ev = detail["evidence"][0]
    assert ev["page_number"] == 2 and "30% by 2030" in ev["evidence_text"] and ev["verified"] is True
    assert detail["conditions"]["quantitative"]["status"] == "met"
    assert detail["conditions"]["benchmark"]["status"] == "not_found"
    assert detail["level_definitions"]["5"].startswith("There are formalized goals")

    summary = client.get(f"/api/assessments/{a['id']}").json()
    assert summary["sustainability_report_found"] and summary["sources_reviewed"] == 1
    assert summary["total"] == 16 and summary["reviewed"] == 0
    assert summary["data_gaps"] == sum(r["evidence_status"] == "NO_EVIDENCE_FOUND" for r in m.values())


def test_full_pipeline_with_mocked_llm_and_review_and_export(client, sample_pdf):
    llm = FakeLLM()
    set_llm_provider(llm)
    a = _create(client)
    _upload(client, a["id"], sample_pdf)
    _run(client, a["id"])
    assert llm.calls, "LLM path should be exercised"
    assert max(len(c) for c in llm.calls) < 6000  # small contexts only

    m = _matrix(client, a["id"])
    planning = m[("energy_management", "planning")]
    assert planning["recommended_level"] == 5 and planning["evidence_status"] == "FOUND"
    assert m[("reporting", "reporting")]["recommended_level"] == 6
    # technical extraction metadata is retained for audit (evidence endpoint) ...
    ev = client.get(f"/api/assessments/{a['id']}/evidence", params={"factor": "energy_management", "dimension": "planning"}).json()
    assert any(e["extraction_method"] == "llm" and e["model_name"] == "fake:test-1b" for e in ev)
    # ... but not exposed in the analyst-facing detail
    d = client.get(f"/api/recommendations/{planning['id']}").json()
    assert "fake:test-1b" not in json.dumps(d) and "triggered_rule" not in d and "confidence" not in d
    assert d["confidence_label"] in ("High", "Medium", "Low") and d["review_state"] == "review"

    # ---- analyst review
    rid = planning["id"]
    bad = client.post(f"/api/recommendations/{rid}/review", json={"action": "OVERRIDE", "selected_level": 4})
    assert bad.status_code == 422  # comment required
    bad = client.post(f"/api/recommendations/{rid}/review", json={"action": "OVERRIDE", "selected_level": 7, "comment": "x"})
    assert bad.status_code == 422
    r = client.post(f"/api/recommendations/{rid}/review",
                    json={"action": "OVERRIDE", "selected_level": 4, "comment": "Target not board-approved.", "analyst": "analyst1"})
    assert r.status_code == 200 and r.json()["review_status"] == "OVERRIDDEN" and r.json()["final_level"] == 4
    assert r.json()["review_state"] == "changed" and r.json()["analyst_comment"] == "Target not board-approved."
    assert r.json()["recommended_level"] == 5  # the original recommendation is kept next to the analyst's answer
    ok = client.post(f"/api/recommendations/{m[('reporting', 'reporting')]['id']}/review", json={"action": "APPROVE"})
    assert ok.json()["review_status"] == "APPROVED" and ok.json()["final_level"] == 6 and ok.json()["review_state"] == "approved"
    detail = client.get(f"/api/recommendations/{rid}").json()
    assert detail["decisions"][0]["previous_level"] == 5 and detail["decisions"][0]["comment"] == "Target not board-approved."

    events = {e["event_type"] for e in client.get(f"/api/assessments/{a['id']}/audit").json()}
    assert {"assessment_created", "document_uploaded", "evidence_extracted", "recommendation_generated",
            "recommendation_overridden", "recommendation_approved"} <= events

    # ---- export
    j = client.get(f"/api/assessments/{a['id']}/export", params={"format": "json"})
    data = json.loads(j.content)
    assert data["assessment"]["company"] == "Nordvik Components AB" and len(data["recommendations"]) == 16
    dec = data["analyst_decisions"][0]
    assert dec["decision"] == "Changed" and dec["ai_recommended_level"] == 5 and dec["analyst_final_level"] == 4
    assert data["data_gaps"] and data["evidence"] and data["sources"]
    assert "fake:test-1b" not in j.text and "llm" not in j.text.lower() and "triggered_rule" not in j.text
    csv_text = client.get(f"/api/assessments/{a['id']}/export", params={"format": "csv"}).content.decode("utf-8-sig")
    assert "Energy Management" in csv_text and "Changed by analyst" in csv_text and "fake:test-1b" not in csv_text
    x = client.get(f"/api/assessments/{a['id']}/export", params={"format": "xlsx"})
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Assessment", "Evidence", "Data Gaps", "Sources", "Analyst Decisions"]
    assert not any("fake:test-1b" in str(c.value) for ws in wb for row in ws.iter_rows() for c in row)
    assert client.get(f"/api/assessments/{a['id']}/export", params={"format": "exe"}).status_code == 422
    assert "export_generated" in {e["event_type"] for e in client.get(f"/api/assessments/{a['id']}/audit").json()}

    # ---- sources
    src = client.get(f"/api/assessments/{a['id']}/sources").json()[0]
    detail = client.get(f"/api/sources/{src['id']}").json()
    assert detail["has_file"] and detail["evidence_count"] > 0 and detail["page_count"] == 5
    f = client.get(f"/api/sources/{src['id']}/file")
    assert f.status_code == 200 and f.content.startswith(b"%PDF")

    # ---- a completed assessment is never silently re-analysed in place
    assert client.post(f"/api/assessments/{a['id']}/run").status_code == 409
    # ---- "Run new analysis" creates a new version that reuses the upload; old decisions are kept
    set_llm_provider(None)
    new = client.post(f"/api/assessments/{a['id']}/rerun")
    assert new.status_code == 202, new.text
    nid = new.json()["id"]
    assert client.get(f"/api/assessments/{nid}/status").json()["status"] == "COMPLETED"
    assert client.get(f"/api/assessments/{nid}/sources").json()[0]["source_type"] == "uploaded_report"
    assert all(r["review_status"] == "PENDING" for r in client.get(f"/api/assessments/{nid}/recommendations").json())
    old = client.get(f"/api/recommendations/{rid}").json()
    assert old["review_status"] == "OVERRIDDEN" and old["decisions"]
