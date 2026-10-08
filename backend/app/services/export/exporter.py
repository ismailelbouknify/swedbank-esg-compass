"""Assessment export: JSON, CSV and Excel (openpyxl)."""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime, timezone
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config.loader import dimension_config, factor_label, source_priority
from app.models.entities import AnalystDecision, Assessment, Evidence, Recommendation, Source
from app.services.presentation import (
    PRODUCT_NAME, PRODUCT_SUBTITLE, REVIEW_STATE_LABEL, SOURCE_STATUS_LABEL, analyst_notices, review_state,
)


def _iso(d: datetime | None) -> str | None:
    return d.isoformat() if d else None


def _level_text(dim: str, level: int | None) -> str | None:
    return dimension_config(dim)["levels"].get(str(level)) if level else None


def build_export(db: Session, a: Assessment) -> dict[str, Any]:
    """Analyst export: company, year, scores, evidence, sources, analyst decisions and data gaps.
    Technical metadata (model, extraction method, rule ids, internal scores) stays in the database."""
    recs = db.scalars(select(Recommendation).where(Recommendation.assessment_id == a.id).order_by(Recommendation.id)).all()
    evidence = db.scalars(select(Evidence).where(Evidence.assessment_id == a.id).order_by(Evidence.id)).all()
    sources = db.scalars(select(Source).where(Source.assessment_id == a.id).order_by(Source.priority, Source.id)).all()
    decisions = db.scalars(select(AnalystDecision).where(AnalystDecision.assessment_id == a.id).order_by(AnalystDecision.id)).all()
    rec_by_id = {r.id: r for r in recs}
    latest_comment: dict[int, str | None] = {d.recommendation_id: d.comment for d in decisions}
    processed = [s for s in sources if s.status == "PROCESSED"]

    def factor_name(key: str) -> str:
        return "Reporting" if key == "reporting" else factor_label(key)

    def question(r: Recommendation) -> str:
        dim = dimension_config(r.dimension)["label"]
        return "Reporting" if r.dimension == "reporting" else f"{factor_label(r.factor)} - {dim}"

    def state(r: Recommendation) -> str:
        return REVIEW_STATE_LABEL[review_state(r.review_status, r.evidence_status)]

    return {
        "report": {"title": PRODUCT_NAME, "subtitle": PRODUCT_SUBTITLE},
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "assessment": {
            "company": a.company.name, "organisation_number": a.company.organisation_number, "website": a.company.website,
            "country": a.company.country, "industry": a.company.industry, "year": a.year,
            "status": a.status.title(), "last_analysed": _iso(a.completed_at), "analysis_duration_seconds": a.duration_seconds,
            "sources_reviewed": len(processed),
            "reviewed": sum(r.review_status != "PENDING" for r in recs), "questions": len(recs),
            "notices": analyst_notices(a.warnings),
        },
        "recommendations": [
            {
                "factor": factor_name(r.factor),
                "dimension": dimension_config(r.dimension)["label"],
                "recommended_level": None if r.evidence_status == "NO_EVIDENCE_FOUND" else r.recommended_level,
                "final_level": r.final_level, "max_level": r.max_level,
                "review_status": state(r),
                "level_description": _level_text(r.dimension, r.final_level or r.recommended_level),
                "reason": r.reason, "confidence": r.confidence_label, "missing_for_next_level": r.next_level_missing,
                "analyst_comment": latest_comment.get(r.id) if r.review_status != "PENDING" else None,
                "evidence_ids": r.evidence_ids,
            }
            for r in recs
        ],
        "evidence": [
            {
                "id": e.id, "factor": factor_name(e.factor), "dimension": dimension_config(e.dimension)["label"],
                "finding": e.claim_text, "quote": e.evidence_text, "source": e.source_title, "page": e.page_number,
                "url": e.source_url, "quote_confirmed_in_source": e.verification_status == "VERIFIED",
            }
            for e in evidence
        ],
        "data_gaps": [
            {"question": question(r), "finding": r.reason, "review_status": state(r)}
            for r in recs if r.evidence_status == "NO_EVIDENCE_FOUND"
        ],
        "sources": [
            {
                "title": s.title or s.original_filename, "type": source_priority(s.source_type)["label"], "url": s.url,
                "file": s.original_filename, "public": s.is_public,
                "status": SOURCE_STATUS_LABEL.get(s.status, s.status.title()),
            }
            for s in sources
        ],
        "analyst_decisions": [
            {
                "question": question(rec_by_id[d.recommendation_id]) if d.recommendation_id in rec_by_id else None,
                "decision": "Approved" if d.action == "APPROVE" else "Changed",
                "ai_recommended_level": rec_by_id[d.recommendation_id].recommended_level if d.recommendation_id in rec_by_id else None,
                "previous_level": d.previous_level, "analyst_final_level": d.selected_level, "comment": d.comment,
                "analyst": d.analyst, "date": _iso(d.created_at),
            }
            for d in decisions
        ],
    }


def to_json(data: dict[str, Any]) -> bytes:
    return json.dumps(data, indent=2, ensure_ascii=False, default=str).encode("utf-8")


RECOMMENDATION_COLUMNS = ["factor", "dimension", "recommended_level", "final_level", "max_level", "review_status",
                          "reason", "level_description", "confidence", "missing_for_next_level", "analyst_comment"]


def _cell(v: Any) -> Any:
    if isinstance(v, (list, dict)):
        return "; ".join(map(str, v)) if isinstance(v, list) else json.dumps(v, ensure_ascii=False)
    return v


def to_csv(data: dict[str, Any]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    a = data["assessment"]
    w.writerow(["company", "year"] + RECOMMENDATION_COLUMNS)
    for r in data["recommendations"]:
        w.writerow([a["company"], a["year"]] + [_cell(r.get(c)) for c in RECOMMENDATION_COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


def to_xlsx(data: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Assessment"
    wb.properties.title = f"{PRODUCT_NAME} - {data['assessment']['company']} {data['assessment']['year']}"
    wb.properties.subject = PRODUCT_SUBTITLE
    a = data["assessment"]
    ws.append([PRODUCT_NAME])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.append([PRODUCT_SUBTITLE])
    ws.append([])
    for k in ("company", "year", "organisation_number", "website", "country", "industry", "status", "last_analysed",
              "sources_reviewed", "reviewed", "questions"):
        ws.append([k, _cell(a.get(k))])
    ws.append([])
    ws.append(RECOMMENDATION_COLUMNS)
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for r in data["recommendations"]:
        ws.append([_cell(r.get(c)) for c in RECOMMENDATION_COLUMNS])

    def sheet(title: str, rows: list[dict[str, Any]], cols: list[str]) -> None:
        s = wb.create_sheet(title)
        s.append(cols)
        for c in s[1]:
            c.font = Font(bold=True)
        for row in rows:
            s.append([_cell(row.get(c)) for c in cols])

    sheet("Evidence", data["evidence"], ["id", "factor", "dimension", "finding", "quote", "source", "page", "url",
                                         "quote_confirmed_in_source"])
    sheet("Data Gaps", data["data_gaps"], ["question", "finding", "review_status"])
    sheet("Sources", data["sources"], ["title", "type", "url", "file", "public", "status"])
    sheet("Analyst Decisions", data["analyst_decisions"], ["question", "decision", "ai_recommended_level", "previous_level",
                                                           "analyst_final_level", "comment", "analyst", "date"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
